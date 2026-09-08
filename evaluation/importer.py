"""Safe XLSX import and blank template generation for SQL evaluations."""

from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from evaluation.contracts import ImportedCase

_REQUIRED_HEADERS = ("需求原文", "真实SQL")
_MAX_CASES = 1_000
_MAX_REQUIREMENT_LENGTH = 10_000
_MAX_SQL_LENGTH = 100_000


class WorkbookValidationError(ValueError):
    """A stable validation error that never embeds workbook cell content."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        row_number: int | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.row_number = row_number


def _error(code: str, message: str, *, row_number: int | None = None) -> None:
    raise WorkbookValidationError(code, message, row_number=row_number)


def _sheet(workbook):
    if "测试案例" in workbook.sheetnames:
        return workbook["测试案例"]
    if len(workbook.sheetnames) == 1:
        return workbook[workbook.sheetnames[0]]
    _error("case_sheet_required", "存在多个工作表时必须包含“测试案例”工作表")


def _header_indexes(cells: tuple[Cell, ...]) -> dict[str, int]:
    values = [str(cell.value).strip() if cell.value is not None else "" for cell in cells]
    if any(values.count(header) > 1 for header in _REQUIRED_HEADERS):
        _error("duplicate_headers", "必填表头不能重复")
    if any(header not in values for header in _REQUIRED_HEADERS):
        _error("missing_headers", "缺少“需求原文”或“真实SQL”表头")
    return {header: values.index(header) for header in _REQUIRED_HEADERS}


def _cell_text(cell: Cell, *, row_number: int) -> str:
    if cell.data_type == "f":
        _error(
            "formula_not_allowed",
            f"第 {row_number} 行必填项不能使用公式",
            row_number=row_number,
        )
    if cell.value is None:
        return ""
    if not isinstance(cell.value, str):
        _error(
            "invalid_cell_type",
            f"第 {row_number} 行必填项必须是文本",
            row_number=row_number,
        )
    return cell.value.strip()


def import_evaluation_workbook(
    content: bytes,
    *,
    max_bytes: int,
) -> tuple[ImportedCase, ...]:
    """Read a two-column evaluation workbook without executing formulas."""

    if len(content) > max_bytes:
        _error("file_too_large", "评测文件超过允许大小")
    if not content:
        _error("invalid_workbook", "评测文件不是有效的 XLSX")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        _error("invalid_workbook", "评测文件不是有效的 XLSX")

    try:
        worksheet = _sheet(workbook)
        rows = worksheet.iter_rows()
        header_cells = next(rows, None)
        if header_cells is None:
            _error("missing_headers", "缺少“需求原文”或“真实SQL”表头")
        indexes = _header_indexes(header_cells)
        cases: list[ImportedCase] = []
        for row_number, cells in enumerate(rows, start=2):
            requirement_cell = cells[indexes["需求原文"]]
            sql_cell = cells[indexes["真实SQL"]]
            requirement = _cell_text(requirement_cell, row_number=row_number)
            reference_sql = _cell_text(sql_cell, row_number=row_number)
            if not requirement and not reference_sql:
                continue
            if not requirement or not reference_sql:
                _error(
                    "missing_required_value",
                    f"第 {row_number} 行缺少必填值",
                    row_number=row_number,
                )
            if len(requirement) > _MAX_REQUIREMENT_LENGTH:
                _error(
                    "requirement_too_long",
                    f"第 {row_number} 行需求原文过长",
                    row_number=row_number,
                )
            if len(reference_sql) > _MAX_SQL_LENGTH:
                _error(
                    "sql_too_long",
                    f"第 {row_number} 行真实 SQL 过长",
                    row_number=row_number,
                )
            cases.append(
                ImportedCase(
                    row_number=row_number,
                    requirement=requirement,
                    reference_sql=reference_sql,
                )
            )
            if len(cases) > _MAX_CASES:
                _error("too_many_cases", "单次评测最多允许 1000 条案例")
        if not cases:
            _error("empty_case_set", "评测文件至少需要一条完整案例")
        return tuple(cases)
    finally:
        workbook.close()


def build_evaluation_template() -> bytes:
    """Build a blank template without embedding business examples."""

    workbook = Workbook()
    cases = workbook.active
    cases.title = "测试案例"
    cases.append(_REQUIRED_HEADERS)
    for _ in range(5):
        cases.append((None, None))
    cases.freeze_panes = "A2"
    cases.auto_filter.ref = "A1:B6"
    cases.column_dimensions["A"].width = 52
    cases.column_dimensions["B"].width = 90
    header_fill = PatternFill("solid", fgColor="173F3A")
    for cell in cases[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in cases.iter_rows(min_row=2, max_row=6, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    cases.row_dimensions[1].height = 26

    instructions = workbook.create_sheet("填写说明")
    instruction_rows = (
        ("项目", "填写要求"),
        ("填写范围", "只填写“需求原文”和“真实SQL”两列。"),
        ("案例粒度", "一行代表一条独立需求，两列必须同时填写。"),
        ("真实 SQL", "填写人工确认的最终 SQL，系统首期不会执行任何 SQL。"),
        ("基准时间", "评测基准时间在评测中心创建任务时统一设置。"),
        ("安全提示", "请仅在本地受控环境保存和上传业务测试集。"),
    )
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 82
    for cell in instructions[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
