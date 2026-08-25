"""Safe, deterministic parsers for metadata uploads."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections.abc import Iterable
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from ontology_core.errors import OntologyImportError
from ontology_core.management.models import DraftDiagnostic, DraftField, DraftObject, UploadLimits
from ontology_core.management.validation import stable_diagnostic_id
from ontology_core.models import FrozenModel
from ontology_core.normalization import normalize_text
from ontology_core.tabular_metadata import (
    ImportDiagnostic,
    TabularMetadataDraft,
    analyze_metadata,
    parse_tabular_objects,
)

_MAX_COMPRESSION_RATIO = 100
_DRIVE_PREFIX = re.compile(r"^[A-Za-z]:")
_TYPE_MAP = {
    "bigint": "integer",
    "bool": "boolean",
    "boolean": "boolean",
    "date": "date",
    "datetime": "dateTime",
    "decimal": "decimal",
    "double": "double",
    "float": "float",
    "int": "integer",
    "integer": "integer",
    "long": "integer",
    "number": "decimal",
    "short": "integer",
    "string": "string",
    "timestamp": "dateTime",
}


class UnsafeUploadError(OntologyImportError):
    """Raised before untrusted upload content reaches a parser library."""


class ParsedUpload(FrozenModel):
    """An immutable, non-persisted parsing result for one uploaded file."""

    file_name: str
    sha256: str
    objects: tuple[DraftObject, ...]
    diagnostics: tuple[DraftDiagnostic, ...]


def parse_metadata_upload(file_name: str, content: bytes, limits: UploadLimits) -> ParsedUpload:
    """Validate and parse a TSV/TXT or macro-free XLSX metadata upload."""
    suffix = Path(file_name).suffix.casefold()
    if not content or len(content) > limits.max_upload_bytes:
        raise UnsafeUploadError("上传文件为空或超过大小限制")
    if suffix in {".tsv", ".txt"}:
        try:
            drafts = parse_tabular_objects(content.decode("utf-8", errors="strict"), file_name)
        except UnicodeDecodeError as exc:
            raise UnsafeUploadError("TSV 文件必须使用 UTF-8 编码") from exc
    elif suffix == ".xlsx":
        _inspect_xlsx_archive(content, limits)
        drafts = _parse_xlsx(content, file_name)
    else:
        raise UnsafeUploadError("仅支持 TSV、TXT 和无宏 XLSX")
    return _to_parsed_upload(file_name, hashlib.sha256(content).hexdigest(), drafts)


def _inspect_xlsx_archive(content: bytes, limits: UploadLimits) -> None:
    """Reject dangerous XLSX archives before openpyxl reads their contents."""
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            expanded_size = 0
            for item in archive.infolist():
                normalized_name = item.filename.replace("\\", "/")
                name = normalized_name.casefold()
                path = PurePosixPath(normalized_name)
                path_parts = normalized_name.split("/")
                if item.flag_bits & 0x1:
                    raise UnsafeUploadError("XLSX 不支持加密条目")
                if (
                    not normalized_name
                    or normalized_name.startswith("/")
                    or _DRIVE_PREFIX.match(normalized_name)
                    or path.is_absolute()
                    or any(part in {".", ".."} for part in path_parts)
                ):
                    raise UnsafeUploadError("XLSX 包含不安全的归档路径")
                if name == "xl/vbaproject.bin" or name.endswith(".vba"):
                    raise UnsafeUploadError("XLSX 不支持宏内容")
                if name.startswith("xl/externallinks/"):
                    raise UnsafeUploadError("XLSX 不支持外部链接")
                if item.file_size and not item.compress_size:
                    raise UnsafeUploadError("XLSX 包含可疑的压缩条目")
                expanded_size += item.file_size
                if expanded_size > limits.max_xlsx_uncompressed_bytes:
                    raise UnsafeUploadError("XLSX 解压后内容超过大小限制")
                if (
                    item.compress_size
                    and item.file_size / item.compress_size > _MAX_COMPRESSION_RATIO
                ):
                    raise UnsafeUploadError("XLSX 条目压缩比过高")
    except (BadZipFile, OSError) as exc:
        raise UnsafeUploadError("XLSX 文件不是有效的 ZIP 归档") from exc


def _parse_xlsx(content: bytes, file_name: str) -> tuple[TabularMetadataDraft, ...]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        drafts: list[TabularMetadataDraft] = []
        for worksheet in workbook.worksheets:
            text = _rows_to_tsv(worksheet.iter_rows(values_only=True))
            if not text.strip():
                continue
            drafts.extend(
                parse_tabular_objects(
                    text,
                    f"{file_name}:{worksheet.title}",
                )
            )
        if not drafts:
            raise OntologyImportError("未解析到有效对象元数据")
        return _merge_drafts(drafts)
    finally:
        workbook.close()


def _rows_to_tsv(rows: Iterable[tuple[object, ...]]) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter="\t", lineterminator="\n")
    for row in rows:
        writer.writerow("" if value is None else str(value) for value in row)
    return output.getvalue()


def _merge_drafts(drafts: Iterable[TabularMetadataDraft]) -> tuple[TabularMetadataDraft, ...]:
    """Merge same-named objects from separate XLSX sheets deterministically."""
    merged: dict[str, TabularMetadataDraft] = {}
    for draft in drafts:
        key = draft.table.physical_name.casefold()
        existing = merged.get(key)
        if existing is None:
            merged[key] = draft
            continue
        table = existing.table.model_copy(
            update={
                "label": existing.table.label or draft.table.label,
                "description": existing.table.description or draft.table.description,
            }
        )
        merged[key] = existing.model_copy(
            update={
                "table": table,
                "fields": (*existing.fields, *draft.fields),
                "object_names": tuple(dict.fromkeys((*existing.object_names, *draft.object_names))),
                "diagnostics": (*existing.diagnostics, *draft.diagnostics),
                "document_diagnostics": (
                    *existing.document_diagnostics,
                    *draft.document_diagnostics,
                ),
                "source_name": "; ".join(dict.fromkeys((existing.source_name, draft.source_name))),
            }
        )
    return tuple(
        sorted(
            merged.values(),
            key=lambda item: (item.table.physical_name.casefold(), item.table.physical_name),
        )
    )


def _to_parsed_upload(
    file_name: str,
    sha256: str,
    drafts: Iterable[TabularMetadataDraft],
) -> ParsedUpload:
    objects: list[DraftObject] = []
    diagnostics: list[DraftDiagnostic] = []
    for draft in drafts:
        object_id = _object_id(draft.table.physical_name)
        metadata_diagnostics = analyze_metadata(draft)
        seen_fields: set[str] = set()
        unique_source_fields = []
        for field in draft.fields:
            field_key = field.physical_name.casefold()
            if field_key not in seen_fields:
                unique_source_fields.append(field)
                seen_fields.add(field_key)
        fields = tuple(
            DraftField(
                id=_field_id(draft.table.physical_name, field.physical_name),
                physical_name=field.physical_name,
                label=field.label,
                description=field.description,
                xsd_type=_TYPE_MAP.get((field.source_type or "").casefold(), "string"),
                primary_key=field.primary_key,
                title=field.title,
                aliases=field.aliases,
            )
            for field in unique_source_fields
        )
        objects.append(
            DraftObject(
                id=object_id,
                physical_name=draft.table.physical_name,
                label=draft.table.label,
                description=draft.table.description,
                fields=fields,
                status="active" if draft.table.enabled else "inactive",
            )
        )
        diagnostics.extend(
            _draft_diagnostic(item, object_id, draft.table.physical_name, draft.source_name)
            for item in metadata_diagnostics
        )
        diagnostics.extend(
            _draft_diagnostic(item, None, None, draft.source_name)
            for item in draft.document_diagnostics
        )
    return ParsedUpload(
        file_name=file_name,
        sha256=sha256,
        objects=tuple(
            sorted(objects, key=lambda item: (item.physical_name.casefold(), item.physical_name))
        ),
        diagnostics=tuple(sorted(diagnostics, key=lambda item: item.id)),
    )


def _draft_diagnostic(
    diagnostic: ImportDiagnostic,
    object_id: str | None,
    table_name: str | None,
    source_name: str,
) -> DraftDiagnostic:
    related_ids = [object_id] if object_id is not None else []
    if diagnostic.field_name is not None and table_name is not None:
        related_ids.append(_field_id(table_name, diagnostic.field_name))
    location = diagnostic.location
    message = diagnostic.message
    if location is not None:
        origin = f"{source_name} " if source_name else ""
        message = f"{message}（{origin}第 {location.line} 行，列“{location.column}”）"
    return DraftDiagnostic(
        id=stable_diagnostic_id(diagnostic.code, tuple(related_ids)),
        code=diagnostic.code,
        severity=diagnostic.severity.value,
        message=message,
        related_ids=tuple(related_ids),
    )


def _object_id(physical_name: str) -> str:
    return f"object/{normalize_text(physical_name)}"


def _field_id(table_name: str, field_name: str) -> str:
    return f"field/{normalize_text(table_name)}/{normalize_text(field_name)}"
