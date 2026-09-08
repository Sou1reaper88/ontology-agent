"""Generate safe, synthetic XLSX samples for ontology metadata imports."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ontology_core.errors import OntologyError
from ontology_core.management.import_format import STANDARD_COLUMNS

_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_REQUIRED_COLUMNS = {
    "对象英文名称",
    "对象中文名称",
    "状态",
    "属性英文名",
    "属性中文名",
    "属性类型",
    "是否主键",
    "是否标题",
}
_COLUMN_WIDTHS = (24, 18, 28, 12, 22, 18, 14, 28, 12, 12)


class ImportTemplateVariantError(OntologyError):
    """Reject unknown template variants through the stable API envelope."""

    code = "import_template_variant_invalid"
    status_code = 400

    def __init__(self) -> None:
        super().__init__("导入样例类型无效")


@dataclass(frozen=True)
class GeneratedImportTemplate:
    """In-memory attachment returned by the management boundary."""

    file_name: str
    media_type: str
    content: bytes


def build_import_template(variant: str) -> GeneratedImportTemplate:
    """Build one macro-free workbook containing only generic placeholders."""
    object_count = {"single": 1, "multiple": 2}.get(variant)
    if object_count is None:
        raise ImportTemplateVariantError()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "对象字段"
    _write_layout(worksheet, object_count)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return GeneratedImportTemplate(
        file_name=f"ontology-import-{variant}.xlsx",
        media_type=_MEDIA_TYPE,
        content=output.getvalue(),
    )


def _write_layout(worksheet, object_count: int) -> None:
    worksheet.merge_cells("A1:J1")
    worksheet.merge_cells("A2:D2")
    worksheet.merge_cells("E2:J2")
    worksheet["A1"] = (
        "注：红色列为必填列；十列名称不可修改。请将虚拟样例全部替换为实际元数据后再上传。"
    )
    worksheet["A2"] = "对象元数据信息"
    worksheet["E2"] = "对象属性信息"
    worksheet.append(STANDARD_COLUMNS)
    for index in range(object_count):
        worksheet.append(_primary_sample_row(index))
        worksheet.append(_secondary_sample_row(index))

    worksheet.freeze_panes = "A4"
    worksheet.sheet_view.showGridLines = False
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[worksheet.cell(row=3, column=index).column_letter].width = width
    _style_layout(worksheet, 3 + object_count * 2)


def _primary_sample_row(index: int) -> tuple[str, ...]:
    suffix = chr(ord("A") + index)
    return (
        f"SAMPLE_OBJECT_{suffix}",
        f"虚拟对象{suffix}",
        "用于演示填写格式的虚拟对象",
        "启用",
        "FIELD_ID",
        "虚拟标识",
        "string",
        "虚拟对象的稳定标识",
        "是",
        "是",
    )


def _secondary_sample_row(index: int) -> tuple[str, ...]:
    suffix = chr(ord("A") + index)
    return (
        f"SAMPLE_OBJECT_{suffix}",
        "",
        "",
        "",
        "FIELD_NAME",
        "虚拟名称",
        "string",
        "",
        "",
        "",
    )


def _style_layout(worksheet, last_row: int) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    note_fill = PatternFill("solid", fgColor="FFF2CC")
    thin_gray = Side(style="thin", color="D9E2F3")

    worksheet["A1"].fill = note_fill
    worksheet["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet["A1"].font = Font(color="7F6000")
    worksheet.row_dimensions[1].height = 42

    for cell in (worksheet["A2"], worksheet["E2"]):
        cell.fill = dark_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[2].height = 24

    for cell in worksheet[3]:
        cell.fill = light_fill
        cell.font = Font(
            color="FFFF0000" if cell.value in _REQUIRED_COLUMNS else "1F1F1F",
            bold=True,
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    worksheet.row_dimensions[3].height = 30

    for row in worksheet.iter_rows(min_row=4, max_row=last_row, min_col=1, max_col=10):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
