from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from api.routes import evaluations
from auth.jwt import get_current_user
from models import Base, EvaluationCase, EvaluationRun, User
from models.base import get_db


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试案例"
    sheet.append(("需求原文", "真实SQL"))
    sheet.append(("synthetic requirement", "SELECT ID FROM T"))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture()
def api_client(monkeypatch) -> Iterator[tuple[TestClient, sessionmaker, User, list[int]]]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(
        engine,
        tables=[EvaluationRun.__table__, EvaluationCase.__table__],
    )
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    current_user = User(id=1, username="owner", display_name="Owner", role_id=1, status=1)
    launched: list[int] = []

    def database() -> Iterator[Session]:
        with factory() as session:
            yield session

    app = FastAPI()
    app.include_router(evaluations.router)
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[get_current_user] = lambda: current_user
    monkeypatch.setattr(evaluations, "_launch_evaluation", launched.append)
    with TestClient(app) as client:
        yield client, factory, current_user, launched


def import_run(client: TestClient) -> dict:
    response = client.post(
        "/evaluations/import",
        data={"name": "首轮基线", "dialect": "hive", "system_time": "2026-08-24"},
        files={
            "file": (
                "cases.xlsx",
                workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_template_download_contains_blank_two_column_workbook(api_client) -> None:
    client, _, _, _ = api_client

    response = client.get("/evaluations/template")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook["测试案例"]["A1"].value == "需求原文"
    assert workbook["测试案例"]["B1"].value == "真实SQL"
    assert workbook["测试案例"]["A2"].value is None


def test_import_creates_private_pending_run_without_storing_workbook(api_client) -> None:
    client, factory, _, _ = api_client

    payload = import_run(client)

    assert payload["status"] == "pending"
    assert payload["total_cases"] == 1
    with factory() as session:
        run = session.get(EvaluationRun, payload["id"])
        assert run is not None and run.user_id == 1
        assert run.cases[0].source_row == 2
        assert run.cases[0].requirement == "synthetic requirement"


def test_list_omits_requirement_and_sql_but_case_detail_returns_them(api_client) -> None:
    client, _, _, _ = api_client
    run = import_run(client)

    listed = client.get("/evaluations").json()
    cases = client.get(f"/evaluations/{run['id']}/cases").json()
    detail = client.get(
        f"/evaluations/{run['id']}/cases/{cases['items'][0]['id']}"
    ).json()

    assert "reference_sql" not in str(listed)
    assert "reference_sql" not in cases["items"][0]
    assert "requirement" not in cases["items"][0]
    assert detail["reference_sql"] == "SELECT ID FROM T"
    assert detail["requirement"] == "synthetic requirement"


def test_start_is_idempotent_and_launches_only_once(api_client) -> None:
    client, _, _, launched = api_client
    run = import_run(client)

    first = client.post(f"/evaluations/{run['id']}/run")
    second = client.post(f"/evaluations/{run['id']}/run")

    assert first.status_code == 202
    assert second.status_code == 202
    assert launched == [run["id"]]


def test_other_user_receives_not_found_for_read_run_and_delete(api_client) -> None:
    client, _, user, _ = api_client
    run = import_run(client)
    user.id = 2

    assert client.get(f"/evaluations/{run['id']}").status_code == 404
    assert client.post(f"/evaluations/{run['id']}/run").status_code == 404
    assert client.delete(f"/evaluations/{run['id']}").status_code == 404


def test_delete_removes_run_and_cases(api_client) -> None:
    client, factory, _, _ = api_client
    run = import_run(client)

    response = client.delete(f"/evaluations/{run['id']}")

    assert response.status_code == 204
    with factory() as session:
        assert session.get(EvaluationRun, run["id"]) is None
        assert session.query(EvaluationCase).count() == 0


def test_invalid_extension_and_date_are_rejected_safely(api_client) -> None:
    client, _, _, _ = api_client
    invalid_file = client.post(
        "/evaluations/import",
        data={"name": "bad", "dialect": "hive", "system_time": "2026-08-24"},
        files={"file": ("cases.txt", b"sensitive", "text/plain")},
    )
    invalid_date = client.post(
        "/evaluations/import",
        data={"name": "bad", "dialect": "hive", "system_time": "not-a-date"},
        files={"file": ("cases.xlsx", workbook_bytes())},
    )

    assert invalid_file.status_code == 400
    assert invalid_file.json()["detail"]["code"] == "invalid_file_type"
    assert "sensitive" not in invalid_file.text
    assert invalid_date.status_code == 422
