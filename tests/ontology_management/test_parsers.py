from __future__ import annotations

import zipfile
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

import ontology_core.management.parsers as parsers
from ontology_core.errors import OntologyImportError
from ontology_core.management.models import UploadLimits
from ontology_core.management.parsers import UnsafeUploadError, parse_metadata_upload

LIMITS = UploadLimits(max_upload_bytes=1_000_000, max_xlsx_uncompressed_bytes=100_000)
STANDARD_COLUMNS = (
    "对象英文名称",
    "对象中文名称",
    "对象描述",
    "状态",
    "属性英文名",
    "属性中文名",
    "属性类型",
    "属性描述",
    "是否主键",
    "是否标题",
)
THREE_OBJECTS_TSV = (
    "对象英文名称\t对象中文名称\t对象描述\t状态\t属性英文名\t属性中文名\t属性类型\t属性描述\t是否主键\t是否标题\n"
    "DEMO_ACCOUNT_M\t账户月表\t账户快照\t启用\tACCOUNT_ID\t账户编码\tbigint\t账户稳定编码\t是\t是\n"
    "DEMO_CUSTOMER_M\t客户月表\t客户快照\t启用\tCUSTOMER_ID\t客户编码\tbigint\t客户稳定编码\t是\t是\n"
    "DEMO_ORDER_D\t订单日表\t订单流水\t启用\tORDER_ID\t订单编码\tbigint\t订单稳定编码\t是\t是\n"
    "demo_account_m\t\t\t\tOPEN_DATE\t开户日期\tdate\t开户日期\t\t\n"
    "DEMO_CUSTOMER_M\t\t\t\tCUSTOMER_NAME\t客户名称\tstring\t客户显示名称\t\t是\n"
    "DEMO_ORDER_D\t\t\t\tCUSTOMER_ID\t客户编码\tbigint\t下单客户编码\t\t\n"
    "DEMO_ACCOUNT_M\t\t\t\tBALANCE\t账户余额\tdecimal\t日终余额\t\t\n"
    "DEMO_ORDER_D\t\t\t\tORDER_AMOUNT\t订单金额\tdecimal\t订单金额\t\t\n"
)


def metadata_row(
    object_name: str,
    field_name: str,
    *,
    object_label: str = "虚拟对象",
    status: str = "启用",
    field_label: str = "虚拟属性",
    field_type: str = "string",
    field_description: str = "",
) -> str:
    return "\t".join(
        (
            object_name,
            object_label,
            "",
            status,
            field_name,
            field_label,
            field_type,
            field_description,
            "",
            "",
        )
    )


def synthetic_xlsx(*sheets: tuple[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


def test_optional_descriptions_do_not_emit_import_warnings() -> None:
    payload = "\t".join(STANDARD_COLUMNS) + "\n" + metadata_row("DEMO_OBJECT", "OBJECT_ID")
    result = parse_metadata_upload("metadata.tsv", payload.encode("utf-8"), LIMITS)

    assert len(result.objects) == 1
    assert not any(item.severity == "warning" for item in result.diagnostics)


def synthetic_xlsx_with_member(name: str, data: bytes) -> bytes:
    archive_name = name.replace("\\", "/")
    payload = synthetic_xlsx(
        (
            "对象",
            [
                ["对象英文名称", "属性英文名"],
                ["DEMO_OBJECT", "OBJECT_ID"],
            ],
        )
    )
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr(archive_name, data)
    content = output.getvalue()
    if archive_name == name:
        return content
    return content.replace(archive_name.encode(), name.encode())


def synthetic_xlsx_with_encrypted_member() -> bytes:
    payload = bytearray(synthetic_xlsx_with_member("encrypted.xml", b"secret"))
    central_directory = payload.rfind(b"PK\x01\x02")
    assert central_directory >= 0
    flags_offset = central_directory + 8
    flags = int.from_bytes(payload[flags_offset : flags_offset + 2], "little")
    payload[flags_offset : flags_offset + 2] = (flags | 0x1).to_bytes(2, "little")
    return bytes(payload)


def test_tsv_parser_groups_three_objects_atomically() -> None:
    parsed = parse_metadata_upload("objects.tsv", THREE_OBJECTS_TSV.encode("utf-8"), LIMITS)

    assert [item.physical_name for item in parsed.objects] == [
        "DEMO_ACCOUNT_M",
        "DEMO_CUSTOMER_M",
        "DEMO_ORDER_D",
    ]
    assert [item.id for item in parsed.objects] == [
        "object/demo_account_m",
        "object/demo_customer_m",
        "object/demo_order_d",
    ]
    assert sum(len(item.fields) for item in parsed.objects) == 8
    assert parsed.objects[0].fields[1].id == "field/demo_account_m/open_date"
    assert parsed.objects[0].fields[1].physical_name == "OPEN_DATE"
    assert parsed.objects[0].fields[1].description == "开户日期"
    assert not any(item.code == "multiple_objects" for item in parsed.diagnostics)


def test_xlsx_finds_the_standard_header_after_instruction_rows() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [
                ["注：请替换全部虚拟样例"],
                ["对象元数据信息", None, None, None, "对象属性信息"],
                list(STANDARD_COLUMNS),
                [
                    "SAMPLE_OBJECT_A",
                    "虚拟对象A",
                    None,
                    "启用",
                    "FIELD_ID",
                    "虚拟标识",
                    "string",
                    None,
                    "是",
                    "是",
                ],
            ],
        )
    )

    parsed = parse_metadata_upload("objects.xlsx", payload, LIMITS)

    assert [item.physical_name for item in parsed.objects] == ["SAMPLE_OBJECT_A"]
    assert [item.physical_name for item in parsed.objects[0].fields] == ["FIELD_ID"]


def test_xlsx_diagnostic_preserves_the_physical_row_after_header_discovery() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [
                ["说明"],
                ["对象元数据信息", None, None, None, "对象属性信息"],
                list(STANDARD_COLUMNS),
                [
                    "SAMPLE_OBJECT_A",
                    "虚拟对象A",
                    None,
                    "启用",
                    "BAD-NAME",
                    "错误字段",
                    "string",
                    None,
                    "",
                    "",
                ],
            ],
        )
    )

    parsed = parse_metadata_upload("objects.xlsx", payload, LIMITS)

    diagnostic = next(
        item for item in parsed.diagnostics if item.code == "invalid_field_identifier"
    )
    assert "objects.xlsx:对象 第 4 行，列“属性英文名”" in diagnostic.message


def test_tsv_rejects_a_partial_or_renamed_header_contract() -> None:
    payload = "对象英文名称\t属性英文名\nSAMPLE_OBJECT_A\tFIELD_ID\n".encode()

    with pytest.raises(OntologyImportError, match="十列标准表头"):
        parse_metadata_upload("objects.tsv", payload, LIMITS)


def test_tsv_rejects_an_extra_named_column() -> None:
    payload = (
        "\t".join((*STANDARD_COLUMNS, "额外列"))
        + "\n"
        + metadata_row("SAMPLE_OBJECT_A", "FIELD_ID")
        + "\t额外值\n"
    ).encode()

    with pytest.raises(OntologyImportError, match="十列标准表头"):
        parse_metadata_upload("objects.tsv", payload, LIMITS)


def test_xlsx_rejects_multiple_standard_headers_in_the_scan_window() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [
                list(STANDARD_COLUMNS),
                [
                    "SAMPLE_OBJECT_A",
                    "虚拟对象A",
                    None,
                    "启用",
                    "FIELD_ID",
                    "虚拟标识",
                    "string",
                    None,
                    "",
                    "",
                ],
                list(STANDARD_COLUMNS),
            ],
        )
    )

    with pytest.raises(OntologyImportError, match="多个十列标准表头"):
        parse_metadata_upload("objects.xlsx", payload, LIMITS)


def test_xlsx_does_not_scan_for_a_header_after_row_twenty() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [[f"说明 {index}"] for index in range(20)]
            + [list(STANDARD_COLUMNS)]
            + [
                [
                    "SAMPLE_OBJECT_A",
                    "虚拟对象A",
                    None,
                    "启用",
                    "FIELD_ID",
                    "虚拟标识",
                    "string",
                    None,
                    "",
                    "",
                ]
            ],
        )
    )

    with pytest.raises(OntologyImportError, match="前 20 行未找到"):
        parse_metadata_upload("objects.xlsx", payload, LIMITS)


def test_tsv_reports_object_and_field_required_value_gaps() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS) + "\nSAMPLE_OBJECT_A\t\t\t\tFIELD_ID\t\t\t\t\t\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    assert {item.code for item in parsed.diagnostics}.issuperset(
        {
            "missing_object_label",
            "missing_object_status",
            "missing_field_label",
            "missing_field_type",
        }
    )


def test_tsv_rejects_unknown_nonblank_boolean_values() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\nSAMPLE_OBJECT_A\t虚拟对象A\t\t启用\tFIELD_ID\t虚拟标识\tstring\t\t可能\t\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    diagnostic = next(item for item in parsed.diagnostics if item.code == "invalid_boolean_value")
    assert "第 2 行，列“是否主键”" in diagnostic.message


def test_tsv_rejects_unknown_object_status_without_guessing() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\nSAMPLE_OBJECT_A\t虚拟对象A\t\t待观察\tFIELD_ID\t虚拟标识\tstring\t\t\t\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    diagnostic = next(item for item in parsed.diagnostics if item.code == "invalid_object_status")
    assert "第 2 行，列“状态”" in diagnostic.message


def test_tsv_reports_conflicting_object_metadata_instead_of_picking_silently() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\nSAMPLE_OBJECT_A\t虚拟对象A\t\t启用\tFIELD_ID\t虚拟标识\tstring\t\t\t\n"
        + "SAMPLE_OBJECT_A\t另一个名称\t\t停用\tFIELD_NAME\t虚拟名称\tstring\t\t\t\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    assert {item.code for item in parsed.diagnostics}.issuperset(
        {"conflicting_object_label", "conflicting_object_status"}
    )


def test_tsv_parser_preserves_source_row_locations_in_diagnostics() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\n"
        + metadata_row("DEMO_ACCOUNT_M", "BAD-NAME")
        + "\n"
        + metadata_row("DEMO_CUSTOMER_M", "CUSTOMER_ID")
        + "\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    diagnostic = next(
        item for item in parsed.diagnostics if item.code == "invalid_field_identifier"
    )
    assert "objects.tsv 第 2 行，列“属性英文名”" in diagnostic.message
    assert diagnostic.related_ids == ("object/demo_account_m", "field/demo_account_m/bad-name")


def test_diagnostic_identity_ignores_the_mutable_source_location_message() -> None:
    first = parse_metadata_upload(
        "objects.tsv",
        (
            "\t".join(STANDARD_COLUMNS) + "\n" + metadata_row("DEMO_ACCOUNT", "BAD-NAME") + "\n"
        ).encode(),
        LIMITS,
    )
    second = parse_metadata_upload(
        "objects.tsv",
        (
            "\t".join(STANDARD_COLUMNS)
            + "\n"
            + metadata_row("DEMO_OTHER", "OTHER_ID")
            + "\n"
            + metadata_row("DEMO_ACCOUNT", "BAD-NAME")
            + "\n"
        ).encode(),
        LIMITS,
    )

    first_diagnostic = next(
        item for item in first.diagnostics if item.code == "invalid_field_identifier"
    )
    second_diagnostic = next(
        item for item in second.diagnostics if item.code == "invalid_field_identifier"
    )

    assert first_diagnostic.related_ids == second_diagnostic.related_ids
    assert first_diagnostic.message != second_diagnostic.message
    assert first_diagnostic.id == second_diagnostic.id


def test_tsv_parser_keeps_missing_object_name_as_a_global_diagnostic() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\n"
        + metadata_row("", "ORPHAN_FIELD")
        + "\n"
        + metadata_row("DEMO_ACCOUNT_M", "ACCOUNT_ID")
        + "\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    diagnostic = next(item for item in parsed.diagnostics if item.code == "missing_table_name")
    assert "objects.tsv 第 2 行，列“对象英文名称”" in diagnostic.message
    assert diagnostic.related_ids == ()


def test_tsv_parser_rejects_invalid_utf8_and_oversized_content() -> None:
    with pytest.raises(UnsafeUploadError, match="UTF-8"):
        parse_metadata_upload("objects.tsv", b"\xff", LIMITS)
    with pytest.raises(UnsafeUploadError, match="为空或超过大小限制"):
        parse_metadata_upload("objects.tsv", b"x" * (LIMITS.max_upload_bytes + 1), LIMITS)


def test_xlsx_parser_groups_sheets_and_interleaved_objects_with_data_only_values() -> None:
    payload = synthetic_xlsx(
        (
            "账户与客户",
            [
                list(STANDARD_COLUMNS),
                [
                    "DEMO_ACCOUNT_M",
                    "账户",
                    "=1+1",
                    "启用",
                    "ACCOUNT_ID",
                    "账户标识",
                    "bigint",
                    None,
                    "",
                    "",
                ],
                [
                    "DEMO_CUSTOMER_M",
                    "客户",
                    "=1+1",
                    "启用",
                    "CUSTOMER_ID",
                    "客户标识",
                    "bigint",
                    None,
                    "",
                    "",
                ],
                [
                    "DEMO_ACCOUNT_M",
                    "账户",
                    "=1+1",
                    "启用",
                    "BALANCE",
                    "余额",
                    "decimal",
                    None,
                    "",
                    "",
                ],
            ],
        ),
        (
            "订单",
            [
                list(STANDARD_COLUMNS),
                [
                    "DEMO_ORDER_D",
                    "订单",
                    None,
                    "启用",
                    "ORDER_ID",
                    "订单标识",
                    "bigint",
                    None,
                    "",
                    "",
                ],
                [
                    "DEMO_ACCOUNT_M",
                    "账户",
                    None,
                    "启用",
                    "ACCOUNT_STATUS",
                    "账户状态",
                    "string",
                    None,
                    "",
                    "",
                ],
                [
                    "DEMO_ORDER_D",
                    "订单",
                    None,
                    "启用",
                    "ORDER_AMOUNT",
                    "订单金额",
                    "decimal",
                    None,
                    "",
                    "",
                ],
            ],
        ),
    )

    parsed = parse_metadata_upload("objects.xlsx", payload, LIMITS)

    assert [item.physical_name for item in parsed.objects] == [
        "DEMO_ACCOUNT_M",
        "DEMO_CUSTOMER_M",
        "DEMO_ORDER_D",
    ]
    assert [field.physical_name for field in parsed.objects[0].fields] == [
        "ACCOUNT_ID",
        "BALANCE",
        "ACCOUNT_STATUS",
    ]
    assert not any(item.code == "unknown_field_type" for item in parsed.diagnostics)


def test_xlsx_rejects_external_links_before_workbook_load(monkeypatch: pytest.MonkeyPatch) -> None:
    payload = synthetic_xlsx_with_member("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
    monkeypatch.setattr(
        parsers, "load_workbook", lambda *args, **kwargs: pytest.fail("已加载工作簿")
    )

    with pytest.raises(UnsafeUploadError, match="外部链接"):
        parse_metadata_upload("objects.xlsx", payload, LIMITS)


@pytest.mark.parametrize(
    ("member", "message"),
    [
        (r"..\evil.xml", "不安全"),
        (r"xl\externalLinks\externalLink1.xml", "外部链接"),
        (r"xl\vbaProject.bin", "宏"),
        (r"C:\evil.xml", "不安全"),
        ("/evil.xml", "不安全"),
        ("./evil.xml", "不安全"),
        ("nested/../evil.xml", "不安全"),
    ],
    ids=[
        "backslash-traversal",
        "backslash-external-link",
        "backslash-macro",
        "drive",
        "absolute",
        "dot",
        "nested-traversal",
    ],
)
def test_xlsx_rejects_unsafe_backslash_or_dot_member_before_workbook_load(
    monkeypatch: pytest.MonkeyPatch,
    member: str,
    message: str,
) -> None:
    if "\\" in member:
        monkeypatch.setattr(zipfile.os, "sep", "/")
    monkeypatch.setattr(
        parsers, "load_workbook", lambda *args, **kwargs: pytest.fail("已加载工作簿")
    )

    with pytest.raises(UnsafeUploadError, match=message):
        parse_metadata_upload("objects.xlsx", synthetic_xlsx_with_member(member, b"unsafe"), LIMITS)


@pytest.mark.parametrize(
    ("member", "data", "message"),
    [
        ("xl/vbaProject.bin", b"macro", "宏"),
        ("large.xml", b"x" * (LIMITS.max_xlsx_uncompressed_bytes + 1), "解压后"),
        ("compressed.xml", b"x" * 20_000, "压缩比"),
    ],
    ids=["macro", "expanded-size", "compression-ratio"],
)
def test_xlsx_rejects_hostile_archive_members(member: str, data: bytes, message: str) -> None:
    with pytest.raises(UnsafeUploadError, match=message):
        parse_metadata_upload("objects.xlsx", synthetic_xlsx_with_member(member, data), LIMITS)


def test_xlsx_rejects_encrypted_member_before_workbook_load(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        parsers, "load_workbook", lambda *args, **kwargs: pytest.fail("已加载工作簿")
    )

    with pytest.raises(UnsafeUploadError, match="加密"):
        parse_metadata_upload("objects.xlsx", synthetic_xlsx_with_encrypted_member(), LIMITS)


def test_xlsm_is_rejected_even_when_payload_is_a_valid_xlsx() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [["对象英文名称", "属性英文名"], ["DEMO_OBJECT", "OBJECT_ID"]],
        )
    )

    with pytest.raises(UnsafeUploadError, match="仅支持"):
        parse_metadata_upload("objects.xlsm", payload, LIMITS)


def test_tsv_duplicate_fields_return_diagnostic_without_model_validation_error() -> None:
    payload = (
        "\t".join(STANDARD_COLUMNS)
        + "\n"
        + metadata_row("DEMO_ACCOUNT_M", "ACCOUNT_ID", field_description="账户编码")
        + "\n"
        + metadata_row("DEMO_ACCOUNT_M", "account_id", field_description="账户编码副本")
        + "\n"
    ).encode()

    parsed = parse_metadata_upload("objects.tsv", payload, LIMITS)

    assert [field.physical_name for field in parsed.objects[0].fields] == ["ACCOUNT_ID"]
    diagnostic = next(item for item in parsed.diagnostics if item.code == "duplicate_field")
    assert "objects.tsv 第 3 行，列“属性英文名”" in diagnostic.message
    assert diagnostic.related_ids == ("object/demo_account_m", "field/demo_account_m/account_id")


def test_xlsx_duplicate_fields_return_diagnostic_without_model_validation_error() -> None:
    payload = synthetic_xlsx(
        (
            "对象",
            [
                list(STANDARD_COLUMNS),
                [
                    "DEMO_ACCOUNT_M",
                    "账户",
                    None,
                    "启用",
                    "ACCOUNT_ID",
                    "账户编码",
                    "string",
                    "账户编码",
                    "",
                    "",
                ],
                [
                    "DEMO_ACCOUNT_M",
                    "账户",
                    None,
                    "启用",
                    "account_id",
                    "账户编码",
                    "string",
                    "账户编码副本",
                    "",
                    "",
                ],
            ],
        )
    )

    parsed = parse_metadata_upload("objects.xlsx", payload, LIMITS)

    assert [field.physical_name for field in parsed.objects[0].fields] == ["ACCOUNT_ID"]
    diagnostic = next(item for item in parsed.diagnostics if item.code == "duplicate_field")
    assert "objects.xlsx:对象 第 3 行，列“属性英文名”" in diagnostic.message
