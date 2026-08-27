from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from ontology_core.management.import_format import STANDARD_COLUMNS
from ontology_core.management.models import UploadLimits
from ontology_core.management.parsers import parse_metadata_upload
from ontology_core.management.templates import (
    ImportTemplateVariantError,
    build_import_template,
)

LIMITS = UploadLimits(max_upload_bytes=1_000_000, max_xlsx_uncompressed_bytes=100_000)
REQUIRED_COLUMNS = {
    "对象英文名称",
    "对象中文名称",
    "状态",
    "属性英文名",
    "属性中文名",
    "属性类型",
    "是否主键",
    "是否标题",
}


def test_single_template_matches_the_reference_layout_and_round_trips() -> None:
    generated = build_import_template("single")

    workbook = load_workbook(BytesIO(generated.content), data_only=True)
    try:
        sheet = workbook.active
        assert sheet.title == "对象字段"
        assert {str(item) for item in sheet.merged_cells.ranges} == {
            "A1:J1",
            "A2:D2",
            "E2:J2",
        }
        assert tuple(cell.value for cell in sheet[3]) == STANDARD_COLUMNS
        assert sheet.freeze_panes == "A4"
        assert all(sheet.column_dimensions[column].width > 10 for column in "ABCDEFGHIJ")
        assert "替换" in sheet["A1"].value
    finally:
        workbook.close()

    parsed = parse_metadata_upload(generated.file_name, generated.content, LIMITS)
    assert [item.physical_name for item in parsed.objects] == ["SAMPLE_OBJECT_A"]
    assert len(parsed.objects[0].fields) == 2


def test_multiple_template_round_trips_two_synthetic_objects() -> None:
    generated = build_import_template("multiple")

    parsed = parse_metadata_upload(generated.file_name, generated.content, LIMITS)

    assert generated.file_name == "ontology-import-multiple.xlsx"
    assert generated.media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert [item.physical_name for item in parsed.objects] == [
        "SAMPLE_OBJECT_A",
        "SAMPLE_OBJECT_B",
    ]
    assert sum(len(item.fields) for item in parsed.objects) == 4


def test_required_headers_are_red_and_optional_descriptions_are_not() -> None:
    generated = build_import_template("single")
    workbook = load_workbook(BytesIO(generated.content), data_only=True)
    try:
        sheet = workbook.active
        colors = {
            cell.value: cell.font.color.rgb
            for cell in sheet[3]
            if cell.font.color is not None and cell.font.color.type == "rgb"
        }
    finally:
        workbook.close()

    assert {column for column, color in colors.items() if color == "FFFF0000"} == REQUIRED_COLUMNS


def test_unknown_template_variant_fails_closed() -> None:
    with pytest.raises(ImportTemplateVariantError, match="样例类型"):
        build_import_template("unsupported")
