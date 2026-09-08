from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from evaluation.contracts import ImportedCase
from evaluation.importer import (
    WorkbookValidationError,
    build_evaluation_template,
    import_evaluation_workbook,
)


def workbook_bytes(
    rows: list[tuple[object, object]],
    *,
    sheet: str = "测试案例",
    second_sheet: str | None = None,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    if second_sheet:
        workbook.create_sheet(second_sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_complete_two_column_workbook() -> None:
    content = workbook_bytes(
        [("需求原文", "真实SQL"), ("需求一", "SELECT ID FROM T")]
    )

    cases = import_evaluation_workbook(content, max_bytes=1024 * 1024)

    assert cases == (
        ImportedCase(row_number=2, requirement="需求一", reference_sql="SELECT ID FROM T"),
    )


def test_unique_sheet_is_accepted_and_blank_rows_are_ignored() -> None:
    content = workbook_bytes(
        [("需求原文", "真实SQL"), (None, None), (" 需求二 ", " SELECT 2 ")],
        sheet="Sheet1",
    )

    cases = import_evaluation_workbook(content, max_bytes=1024 * 1024)

    assert cases[0].row_number == 3
    assert cases[0].requirement == "需求二"
    assert cases[0].reference_sql == "SELECT 2"


def test_missing_value_reports_row_without_echoing_sensitive_content() -> None:
    content = workbook_bytes(
        [("需求原文", "真实SQL"), ("sensitive requirement", None)]
    )

    with pytest.raises(WorkbookValidationError) as caught:
        import_evaluation_workbook(content, max_bytes=1024 * 1024)

    assert caught.value.code == "missing_required_value"
    assert caught.value.row_number == 2
    assert "sensitive" not in str(caught.value).casefold()


def test_multiple_sheets_without_named_case_sheet_are_rejected() -> None:
    content = workbook_bytes(
        [("需求原文", "真实SQL"), ("需求", "SELECT 1")],
        sheet="Sheet1",
        second_sheet="填写说明",
    )

    with pytest.raises(WorkbookValidationError) as caught:
        import_evaluation_workbook(content, max_bytes=1024 * 1024)

    assert caught.value.code == "case_sheet_required"


def test_formula_in_required_cell_is_rejected_without_calculation() -> None:
    content = workbook_bytes(
        [("需求原文", "真实SQL"), ("需求", "=CONCAT(\"SELECT\",\" 1\")")]
    )

    with pytest.raises(WorkbookValidationError) as caught:
        import_evaluation_workbook(content, max_bytes=1024 * 1024)

    assert caught.value.code == "formula_not_allowed"
    assert caught.value.row_number == 2


def test_invalid_headers_and_empty_case_set_are_rejected() -> None:
    invalid = workbook_bytes([("需求原文", "答案")])
    empty = workbook_bytes([("需求原文", "真实SQL"), (None, None)])

    with pytest.raises(WorkbookValidationError) as invalid_error:
        import_evaluation_workbook(invalid, max_bytes=1024 * 1024)
    with pytest.raises(WorkbookValidationError) as empty_error:
        import_evaluation_workbook(empty, max_bytes=1024 * 1024)

    assert invalid_error.value.code == "missing_headers"
    assert empty_error.value.code == "empty_case_set"


def test_oversized_and_invalid_files_are_rejected() -> None:
    with pytest.raises(WorkbookValidationError) as oversized:
        import_evaluation_workbook(b"x" * 11, max_bytes=10)
    with pytest.raises(WorkbookValidationError) as invalid:
        import_evaluation_workbook(b"not-an-xlsx", max_bytes=100)

    assert oversized.value.code == "file_too_large"
    assert invalid.value.code == "invalid_workbook"


def test_generated_template_has_only_blank_business_input_rows() -> None:
    workbook = load_workbook(BytesIO(build_evaluation_template()), data_only=False)

    assert workbook.sheetnames == ["测试案例", "填写说明"]
    worksheet = workbook["测试案例"]
    assert worksheet["A1"].value == "需求原文"
    assert worksheet["B1"].value == "真实SQL"
    assert all(
        worksheet.cell(row=row, column=column).value is None
        for row in range(2, 7)
        for column in (1, 2)
    )
    instructions = " ".join(
        str(cell.value or "") for row in workbook["填写说明"].iter_rows() for cell in row
    )
    assert "不会执行" in instructions
    assert "基准时间" in instructions
